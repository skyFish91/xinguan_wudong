# -*- coding: utf-8 -*-
"""将合并大纲 Markdown 转为 Word (.docx)。处理：#/##/### 标题、|表格|、- 列表、> 引用、```代码块、**粗体**、`行内代码`、--- 分隔线。"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

CN_FONT = '宋体'
MONO_FONT = 'Consolas'

def set_font(run, size=None, bold=None, mono=False):
    name = MONO_FONT if mono else CN_FONT
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT if not mono else name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold

def add_inline(par, text, size=None):
    for tok in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            r = par.add_run(tok[2:-2]); set_font(r, size=size, bold=True)
        elif tok.startswith('`') and tok.endswith('`'):
            r = par.add_run(tok[1:-1]); set_font(r, size=size, mono=True)
        else:
            r = par.add_run(tok); set_font(r, size=size)

def add_table(doc, rows):
    # rows: list of list[str]，第一行是表头；跳过 |---| 分隔行由调用方保证
    header = [c.strip() for c in rows[0]]
    body = [[c.strip() for c in r] for r in rows[1:]]
    ncol = len(header)
    table = doc.add_table(rows=1, cols=ncol)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    for i, h in enumerate(header):
        p = hdr[i].paragraphs[0]; r = p.add_run(h); set_font(r, size=10.5, bold=True)
    for row in body:
        cells = table.add_row().cells
        for i in range(ncol):
            txt = row[i] if i < len(row) else ''
            p = cells[i].paragraphs[0]; add_inline(p, txt, size=10.5)
    return table

def main(md_path, out_path):
    with open(md_path, encoding='utf-8') as f:
        lines = f.read().splitlines()

    doc = Document()
    # 默认样式
    normal = doc.styles['Normal']
    normal.font.name = CN_FONT
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)
    normal.font.size = Pt(11)

    i = 0
    n = len(lines)
    in_code = False
    code_buf = []
    while i < n:
        line = lines[i]
        s = line.rstrip()

        # 代码块
        if s.strip().startswith('```'):
            if in_code:
                # 结束代码块，输出
                for cl in code_buf:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    r = p.add_run(cl if cl else ' '); set_font(r, size=9, mono=True)
                in_code = False; code_buf = []
            else:
                in_code = True; code_buf = []
            i += 1; continue
        if in_code:
            code_buf.append(s)
            i += 1; continue

        # 空行
        if not s.strip():
            i += 1; continue

        # 标题
        m = re.match(r'^(#{1,6})\s+(.*)$', s)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            h = doc.add_heading(level=min(level, 4))
            h.paragraph_format.space_before = Pt(10 if level <= 2 else 6)
            add_inline(h, text, size={1: 20, 2: 16, 3: 14, 4: 12}.get(level, 12))
            i += 1; continue

        # 表格：以 | 开头的连续行
        if s.strip().startswith('|'):
            tbl = []
            while i < n and lines[i].strip().startswith('|'):
                row = lines[i].strip().strip('|').split('|')
                tbl.append(row)
                i += 1
            # 去除分隔行（全部为 - : 空格）
            tbl = [r for r in tbl if not all(re.fullmatch(r'[\s:—-]*', c) for c in r)]
            if tbl:
                add_table(doc, tbl)
            continue

        # 分隔线
        if re.fullmatch(r'-{3,}', s.strip()):
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
            r = p.add_run('—' * 30); set_font(r, size=9); r.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            i += 1; continue

        # 引用
        if s.strip().startswith('>'):
            text = s.strip().lstrip('>').strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(12)
            add_inline(p, text, size=10)
            for r in p.runs:
                r.font.italic = True
                r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            i += 1; continue

        # 列表
        if re.match(r'^\s*[-*]\s+', s):
            text = re.sub(r'^\s*[-*]\s+', '', s)
            p = doc.add_paragraph(style='List Bullet')
            add_inline(p, text, size=11)
            i += 1; continue

        # 普通段落
        p = doc.add_paragraph()
        add_inline(p, s, size=11)
        i += 1

    doc.save(out_path)
    print('saved:', out_path)

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
